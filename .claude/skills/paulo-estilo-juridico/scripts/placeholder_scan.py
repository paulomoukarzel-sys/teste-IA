#!/usr/bin/env python3
"""
placeholder_scan.py — Varre arquivos .txt e .docx em pastas pipeline/ de
clientes buscando placeholders não resolvidos e gera data/placeholders.json.

Um placeholder é qualquer texto entre colchetes que NÃO seja uma referência
legítima a artigos, parágrafos, incisos, documentos, etc.

Uso:
    python placeholder_scan.py
"""

import sys
import os
import json
import re
from pathlib import Path
from datetime import datetime

sys.stdout.reconfigure(encoding='utf-8')

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent.parent.parent

CLIENTES_DIR = PROJECT_ROOT / "Clientes"

OUTPUT_DIR = PROJECT_ROOT / "data"
OUTPUT_FILE = OUTPUT_DIR / "placeholders.json"

# Referências legítimas — NÃO devem ser tratadas como placeholder
PATTERN_PLACEHOLDER = re.compile(
    r"\[(?!"
    r"arts?\.|art\.|§|inc\.|n\.|p\.|cf\.|vide|v\.|doc\.|event\.|fl\."
    r").*?\]",
    re.IGNORECASE,
)


# ---------------------------------------------------------------------------
# Extratores de texto
# ---------------------------------------------------------------------------

def extrair_txt(caminho: Path) -> list[tuple[int, str]]:
    """Retorna lista de (número_linha, texto_linha) para arquivo .txt."""
    linhas = []
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            with open(caminho, encoding=enc) as f:
                for i, linha in enumerate(f, start=1):
                    linhas.append((i, linha.rstrip("\n")))
            return linhas
        except (UnicodeDecodeError, OSError):
            continue
    return linhas


def extrair_docx(caminho: Path) -> list[tuple[int, str]]:
    """
    Retorna lista de (número_parágrafo, texto_parágrafo) para arquivo .docx.
    Requer python-docx. Se não disponível, ignora silenciosamente.
    """
    try:
        from docx import Document  # type: ignore
    except ImportError:
        return []

    try:
        doc = Document(str(caminho))
        linhas = []
        for i, para in enumerate(doc.paragraphs, start=1):
            texto = para.text.strip()
            if texto:
                linhas.append((i, texto))
        return linhas
    except Exception:
        return []


def extrair_linhas(caminho: Path) -> list[tuple[int, str]]:
    """Despacha para o extrator correto com base na extensão."""
    ext = caminho.suffix.lower()
    if ext == ".txt":
        return extrair_txt(caminho)
    if ext == ".docx":
        return extrair_docx(caminho)
    return []


# ---------------------------------------------------------------------------
# Varredura
# ---------------------------------------------------------------------------

def encontrar_pipeline_dirs(base: Path) -> list[Path]:
    """Localiza todos os diretórios pipeline/ dentro da árvore base."""
    dirs = []
    if not base.exists():
        print(f"  [aviso] Pasta não encontrada: {base}")
        return dirs
    for item in base.rglob("pipeline"):
        if item.is_dir():
            dirs.append(item)
    return dirs


def varrer_pipeline(pipeline_dir: Path) -> list[dict]:
    """Varre um diretório pipeline/ e retorna lista de ocorrências de placeholders."""
    ocorrencias = []
    for ext in ("*.txt", "*.docx"):
        for arquivo in pipeline_dir.glob(ext):
            linhas = extrair_linhas(arquivo)
            for num_linha, texto in linhas:
                for match in PATTERN_PLACEHOLDER.finditer(texto):
                    ocorrencias.append({
                        "arquivo": arquivo.name,
                        "linha": num_linha,
                        "placeholder": match.group(),
                        "_path": str(arquivo),
                    })
    return ocorrencias


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def nome_cliente_de_pipeline(pipeline_dir: Path) -> str:
    """
    Extrai o nome do cliente a partir do caminho do pipeline.
    Exemplo: Clientes/Leonardo - Le Motos/pipeline → "Leonardo - Le Motos"
    """
    # O parent do pipeline/ é a pasta do cliente (ou uma subpasta)
    partes = pipeline_dir.parts
    idx_clientes = None
    for i, parte in enumerate(partes):
        if parte.lower() == "clientes":
            idx_clientes = i
            break
    if idx_clientes is not None and idx_clientes + 1 < len(partes):
        return partes[idx_clientes + 1]
    return pipeline_dir.parent.name


def formatar_tabela(por_caso: dict[str, list]) -> None:
    """Imprime tabela de resultados no terminal."""
    total = sum(len(v) for v in por_caso.values())
    if total == 0:
        print("  Nenhum placeholder encontrado.")
        return

    print(f"\n{'CASO':<35} {'ARQUIVO':<45} {'LINHA':>6}  PLACEHOLDER")
    print("-" * 110)
    for caso in sorted(por_caso):
        for item in por_caso[caso]:
            print(
                f"{caso:<35} {item['arquivo']:<45} {item['linha']:>6}  {item['placeholder']}"
            )
    print("-" * 110)
    print(f"Total: {total} placeholder(s) em {len(por_caso)} caso(s)")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print("=" * 60)
    print("Scanner de Placeholders — pastas pipeline/")
    print(f"Raiz do projeto: {PROJECT_ROOT}")
    print("=" * 60)

    pipeline_dirs = encontrar_pipeline_dirs(CLIENTES_DIR)
    print(f"\nPastas pipeline/ encontradas: {len(pipeline_dirs)}")

    por_caso: dict[str, list] = {}
    total_geral = 0

    for pdir in sorted(pipeline_dirs):
        cliente = nome_cliente_de_pipeline(pdir)
        ocorrencias = varrer_pipeline(pdir)

        if not ocorrencias:
            continue

        # Remover campo interno _path do JSON e ordenar por linha
        ocorrencias_clean = [
            {k: v for k, v in o.items() if k != "_path"}
            for o in sorted(ocorrencias, key=lambda x: (x["arquivo"], x["linha"]))
        ]

        if cliente not in por_caso:
            por_caso[cliente] = []
        por_caso[cliente].extend(ocorrencias_clean)
        total_geral += len(ocorrencias_clean)

    formatar_tabela(por_caso)

    resultado = {
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "total": total_geral,
        "por_caso": dict(sorted(por_caso.items())),
    }

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(resultado, f, ensure_ascii=False, indent=2)

    print(f"\nJSON salvo em: {OUTPUT_FILE}")
    print("=" * 60)


if __name__ == "__main__":
    main()
